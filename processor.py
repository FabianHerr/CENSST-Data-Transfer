import pandas as pd
import openpyxl
import os
import glob

def process_single_file(source_file, period_to_find):
    """
    Process a single Honoraires file and return new rows data.
    Returns: (new_rows_data, warning_message)
    """
    try:
        # Normalize path for Windows (handles spaces and special characters)
        normalized_file = os.path.normpath(source_file)
        
        # Verify file exists before attempting to open
        if not os.path.exists(normalized_file):
            error_msg = f"Le fichier n'existe pas: {normalized_file}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return [], f"Error: File not found - {normalized_file}"
        
        print(f"DEBUG: Tentative d'ouverture de : {normalized_file}")
        
        # Load source file
        df_source = pd.read_excel(normalized_file, sheet_name='CNESST', header=None, engine='openpyxl')
        
        # 1. Extract Audiologist Name from Cell C3 (Row 2, Col 2)
        audiologist = str(df_source.iloc[2, 2]).strip()
        if not audiologist or audiologist.lower() in ['nan', 'none', '']:
            return [], f"Warning: Could not find Audiologist name in {os.path.basename(normalized_file)}"
        
        # 2. Find the period row in Column C (Index 2)
        period_col = df_source.iloc[:, 2].astype(str).str.strip()
        period_rows = df_source[period_col == str(period_to_find)]
        
        if period_rows.empty:
            return [], f"Warning: No data found for Period {period_to_find} in {os.path.basename(normalized_file)}"
        
        # 3. Extract all patient rows following each period header
        new_rows_data = []
        processed_rows = set()
        
        for period_idx, period_row in period_rows.iterrows():
            start_row = period_idx + 1
            
            for row_idx in range(start_row, len(df_source)):
                if row_idx in processed_rows:
                    continue
                
                current_period_value = str(df_source.iloc[row_idx, 2]).strip()
                if current_period_value.isdigit() and current_period_value != str(period_to_find):
                    break
                
                patient_name = str(df_source.iloc[row_idx, 3]).strip()
                if (current_period_value in ['', 'nan', 'None'] and 
                    patient_name in ['', 'nan', 'None', 'Nom du patient']):
                    break
                
                if patient_name.lower() in ['nom du patient', 'patient', 'name', 'nom']:
                    processed_rows.add(row_idx)
                    continue
                
                if not patient_name or patient_name.lower() in ['nan', 'none', '']:
                    processed_rows.add(row_idx)
                    continue
                
                date_value = df_source.iloc[row_idx, 4]
                
                # Create new entry - ensure we have enough columns
                num_cols = max(11, 11)  # At least up to column K
                new_entry = [None] * num_cols
                
                new_entry[1] = patient_name      # Column B: Patient Name
                new_entry[2] = date_value         # Column C: Date
                new_entry[5] = audiologist        # Column F: Audiologist Name
                new_entry[10] = str(period_to_find)  # Column K: Period Number
                
                new_rows_data.append(new_entry)
                processed_rows.add(row_idx)
        
        return new_rows_data, None
        
    except Exception as e:
        error_msg = f"Error processing {os.path.basename(source_file)}: {str(e)}"
        print(f"DEBUG: ERREUR lors du traitement - {error_msg}")
        return [], error_msg

def write_rows_to_sheet(ws, new_rows_data, next_row, source_row):
    """Helper function to write rows to the sheet with formatting"""
    for new_row_data in new_rows_data:
        for col_idx, value in enumerate(new_row_data, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=value)
            
            # Copy formatting from source row if available
            if source_row is not None and col_idx <= len(source_row):
                source_cell = source_row[col_idx - 1]
                # Copy font
                if source_cell.font:
                    cell.font = openpyxl.styles.Font(
                        name=source_cell.font.name,
                        size=source_cell.font.size,
                        bold=source_cell.font.bold,
                        italic=source_cell.font.italic,
                        underline=source_cell.font.underline,
                        strike=source_cell.font.strike,
                        color=source_cell.font.color
                    )
                # Copy fill
                if source_cell.fill:
                    cell.fill = openpyxl.styles.PatternFill(
                        fill_type=source_cell.fill.fill_type,
                        start_color=source_cell.fill.start_color,
                        end_color=source_cell.fill.end_color
                    )
                # Copy border
                if source_cell.border:
                    cell.border = openpyxl.styles.Border(
                        left=source_cell.border.left,
                        right=source_cell.border.right,
                        top=source_cell.border.top,
                        bottom=source_cell.border.bottom
                    )
                # Copy alignment, but always enable text wrapping
                if source_cell.alignment:
                    cell.alignment = openpyxl.styles.Alignment(
                        horizontal=source_cell.alignment.horizontal,
                        vertical=source_cell.alignment.vertical,
                        wrap_text=True,  # Always enable text wrapping
                        shrink_to_fit=source_cell.alignment.shrink_to_fit,
                        indent=source_cell.alignment.indent
                    )
                else:
                    cell.alignment = openpyxl.styles.Alignment(
                        wrap_text=True,
                        vertical='top'
                    )
                # Copy number format
                if source_cell.number_format:
                    cell.number_format = source_cell.number_format
            else:
                # If no source row, set default alignment with text wrapping enabled
                cell.alignment = openpyxl.styles.Alignment(
                    wrap_text=True,
                    vertical='top'
                )
        
        next_row += 1
    
    return next_row

def run_transfer_logic(source_path, target_path, period_to_find):
    """
    Transfer data from multiple Honoraires files in a folder to SUIVI master file.
    
    source_path: Folder containing Honoraires Excel files
    target_path: Path to the SUIVI master file
    
    Source Sheet ("CNESST"):
    - Audiologist Name: Cell C3 (Row 2, Col 2)
    - Period numbers (1-26): Column C (Index 2)
    - Patient Names: Column D (Index 3)
    - Dates: Column E (Index 4)
    
    Destination Sheet ("Suivi"):
    - Column B: Patient Name
    - Column C: Date
    - Column F: Audiologist Name
    - Column K: Period Number
    """
    try:
        # Normalize paths for Windows (handles spaces and special characters)
        normalized_source = os.path.normpath(source_path)
        normalized_target = os.path.normpath(target_path)
        
        print(f"DEBUG: run_transfer_logic - Dossier source normalisé: {normalized_source}")
        print(f"DEBUG: run_transfer_logic - Fichier cible normalisé: {normalized_target}")
        
        # Verify paths exist
        if not os.path.exists(normalized_source):
            error_msg = f"Error: Le dossier source n'existe pas: {normalized_source}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        if not os.path.exists(normalized_target):
            error_msg = f"Error: Le fichier cible n'existe pas: {normalized_target}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Check if source_path is a directory
        if not os.path.isdir(normalized_source):
            error_msg = f"Error: {normalized_source} is not a valid folder."
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Find all Excel files in the folder
        excel_files = []
        for ext in ['*.xlsx', '*.xlsm', '*.XLSX', '*.XLSM']:
            search_pattern = os.path.join(normalized_source, ext)
            print(f"DEBUG: Recherche de fichiers avec le pattern: {search_pattern}")
            found_files = glob.glob(search_pattern)
            excel_files.extend(found_files)
            print(f"DEBUG: Fichiers trouvés avec {ext}: {len(found_files)}")
        
        if not excel_files:
            error_msg = f"Error: No Excel files found in the folder: {normalized_source}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Normalize all file paths
        excel_files = [os.path.normpath(f) for f in excel_files]
        
        # Sort files for consistent processing
        excel_files.sort()
        print(f"DEBUG: Total fichiers Excel trouvés: {len(excel_files)}")
        for idx, f in enumerate(excel_files, 1):
            print(f"DEBUG:   {idx}. {f}")
        
        # Verify target file exists before loading
        if not os.path.exists(normalized_target):
            error_msg = f"Error: Le fichier cible n'existe pas: {normalized_target}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        print(f"DEBUG: Tentative d'ouverture du fichier cible: {normalized_target}")
        # Load workbook once at the beginning
        wb = openpyxl.load_workbook(normalized_target, keep_vba=True)
        
        # Get or create the Suivi sheet
        if 'Suivi' in wb.sheetnames:
            ws = wb['Suivi']
            # Find the last row that has data
            last_row_with_data = 0
            for row_num in range(ws.max_row, 0, -1):
                cell_b = ws.cell(row=row_num, column=2).value
                cell_c = ws.cell(row=row_num, column=3).value
                b_has_data = cell_b is not None and str(cell_b).strip() != '' and str(cell_b).strip().lower() != 'nan'
                c_has_data = cell_c is not None and str(cell_c).strip() != '' and str(cell_c).strip().lower() != 'nan'
                if b_has_data or c_has_data:
                    last_row_with_data = row_num
                    break
            next_row = last_row_with_data + 1 if last_row_with_data > 0 else 1
        else:
            ws = wb.create_sheet('Suivi')
            next_row = 1
        
        # Get formatting template (find once, reuse for all files)
        source_row = None
        if ws.max_row > 0:
            for row_num in range(max(1, next_row - 1), 0, -1):
                cell_b = ws.cell(row=row_num, column=2).value
                cell_c = ws.cell(row=row_num, column=3).value
                if (cell_b is not None and str(cell_b).strip() != '' and str(cell_b).strip().lower() != 'nan') or \
                   (cell_c is not None and str(cell_c).strip() != '' and str(cell_c).strip().lower() != 'nan'):
                    source_row = ws[row_num]
                    break
            if source_row is None and ws.max_row > 0:
                source_row = ws[ws.max_row]
        
        # Process each file
        total_added = 0
        files_processed = 0
        warnings = []
        
        for excel_file in excel_files:
            new_rows_data, warning_msg = process_single_file(
                excel_file, period_to_find
            )
            
            if warning_msg:
                warnings.append(warning_msg)
            
            if new_rows_data:
                # Write rows to sheet
                next_row = write_rows_to_sheet(ws, new_rows_data, next_row, source_row)
                total_added += len(new_rows_data)
                files_processed += 1
                # Save after each file to preserve progress
                print(f"DEBUG: Sauvegarde du fichier cible: {normalized_target}")
                wb.save(normalized_target)
                
                # Rename the file to indicate successful transfer
                try:
                    file_dir = os.path.dirname(excel_file)
                    file_name = os.path.basename(excel_file)
                    new_name = f"(TRANSFER COMPLETED)-{file_name}"
                    new_path = os.path.join(file_dir, new_name)
                    
                    # If a file with the new name already exists, add a number suffix
                    counter = 1
                    original_new_path = new_path
                    while os.path.exists(new_path):
                        name_without_ext, ext = os.path.splitext(original_new_path)
                        new_path = f"{name_without_ext} ({counter}){ext}"
                        counter += 1
                    
                    os.rename(excel_file, new_path)
                except Exception as rename_error:
                    # If renaming fails, add a warning but don't stop processing
                    warnings.append(f"Warning: Could not rename {os.path.basename(excel_file)}: {str(rename_error)}")
        
        # Build result message
        result_parts = [f"Processed {files_processed} file(s). Added {total_added} total records for Period {period_to_find}."]
        if warnings:
            result_parts.append(f"\nWarnings:\n" + "\n".join(warnings))
        
        if total_added == 0:
            return f"No records found for Period {period_to_find} in any files.\n" + "\n".join(warnings) if warnings else ""
        
        return "\n".join(result_parts)

    except FileNotFoundError as e:
        return f"Error: File not found - {str(e)}"
    except ValueError as e:
        if "Worksheet named 'CNESST' not found" in str(e) or "Worksheet named 'Suivi' not found" in str(e):
            return f"Error: Required worksheet not found. Please ensure files contain 'CNESST' and 'Suivi' sheets."
        return f"Error: {str(e)}"
    except Exception as e:
        return f"Error: {str(e)}"

def run_transfer_all_periods(source_path, target_path):
    """
    Transfer data from Honoraires files for ALL periods (1-26) that have data.
    
    Processes each period sequentially and transfers all matching data.
    """
    try:
        # Normalize paths for Windows (handles spaces and special characters)
        normalized_source = os.path.normpath(source_path)
        normalized_target = os.path.normpath(target_path)
        
        print(f"DEBUG: run_transfer_all_periods - Dossier source normalisé: {normalized_source}")
        print(f"DEBUG: run_transfer_all_periods - Fichier cible normalisé: {normalized_target}")
        
        # Verify paths exist
        if not os.path.exists(normalized_source):
            error_msg = f"Error: Le dossier source n'existe pas: {normalized_source}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        if not os.path.exists(normalized_target):
            error_msg = f"Error: Le fichier cible n'existe pas: {normalized_target}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Check if source_path is a directory
        if not os.path.isdir(normalized_source):
            error_msg = f"Error: {normalized_source} is not a valid folder."
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Find all Excel files in the folder (exclude already processed files)
        excel_files = []
        for ext in ['*.xlsx', '*.xlsm', '*.XLSX', '*.XLSM']:
            search_pattern = os.path.join(normalized_source, ext)
            print(f"DEBUG: Recherche de fichiers avec le pattern: {search_pattern}")
            found_files = glob.glob(search_pattern)
            excel_files.extend(found_files)
            print(f"DEBUG: Fichiers trouvés avec {ext}: {len(found_files)}")
        
        # Filter out files that start with "(TRANSFER COMPLETED)-"
        excel_files = [f for f in excel_files if not os.path.basename(f).startswith("(TRANSFER COMPLETED)-")]
        
        if not excel_files:
            error_msg = f"Error: No unprocessed Excel files found in the folder: {normalized_source}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        # Normalize all file paths
        excel_files = [os.path.normpath(f) for f in excel_files]
        
        # Sort files for consistent processing
        excel_files.sort()
        print(f"DEBUG: Total fichiers Excel non traités trouvés: {len(excel_files)}")
        for idx, f in enumerate(excel_files, 1):
            print(f"DEBUG:   {idx}. {f}")
        
        # Verify target file exists before loading
        if not os.path.exists(normalized_target):
            error_msg = f"Error: Le fichier cible n'existe pas: {normalized_target}"
            print(f"DEBUG: ERREUR - {error_msg}")
            return error_msg
        
        print(f"DEBUG: Tentative d'ouverture du fichier cible: {normalized_target}")
        # Load workbook once at the beginning
        wb = openpyxl.load_workbook(normalized_target, keep_vba=True)
        
        # Get or create the Suivi sheet
        if 'Suivi' in wb.sheetnames:
            ws = wb['Suivi']
            # Find the last row that has data
            last_row_with_data = 0
            for row_num in range(ws.max_row, 0, -1):
                cell_b = ws.cell(row=row_num, column=2).value
                cell_c = ws.cell(row=row_num, column=3).value
                b_has_data = cell_b is not None and str(cell_b).strip() != '' and str(cell_b).strip().lower() != 'nan'
                c_has_data = cell_c is not None and str(cell_c).strip() != '' and str(cell_c).strip().lower() != 'nan'
                if b_has_data or c_has_data:
                    last_row_with_data = row_num
                    break
            next_row = last_row_with_data + 1 if last_row_with_data > 0 else 1
        else:
            ws = wb.create_sheet('Suivi')
            next_row = 1
        
        # Get formatting template (find once, reuse for all files)
        source_row = None
        if ws.max_row > 0:
            for row_num in range(max(1, next_row - 1), 0, -1):
                cell_b = ws.cell(row=row_num, column=2).value
                cell_c = ws.cell(row=row_num, column=3).value
                if (cell_b is not None and str(cell_b).strip() != '' and str(cell_b).strip().lower() != 'nan') or \
                   (cell_c is not None and str(cell_c).strip() != '' and str(cell_c).strip().lower() != 'nan'):
                    source_row = ws[row_num]
                    break
            if source_row is None and ws.max_row > 0:
                source_row = ws[ws.max_row]
        
        # Process all periods (1-26)
        total_added_all = 0
        periods_processed = []
        periods_with_data = []
        all_warnings = []
        files_with_data = set()  # Track which files had data transferred
        
        for period in range(1, 27):
            period_str = str(period)
            period_added = 0
            
            # Process each file for this period
            for excel_file in excel_files:
                # Normalize file path
                normalized_excel_file = os.path.normpath(excel_file)
                
                # Skip files that have already been renamed
                if os.path.basename(normalized_excel_file).startswith("(TRANSFER COMPLETED)-"):
                    continue
                
                # Verify file still exists
                if not os.path.exists(normalized_excel_file):
                    print(f"DEBUG: Avertissement - Le fichier n'existe plus: {normalized_excel_file}")
                    continue
                
                new_rows_data, warning_msg = process_single_file(
                    normalized_excel_file, period_str
                )
                
                if warning_msg and "Error" in warning_msg:
                    all_warnings.append(warning_msg)
                
                if new_rows_data:
                    # Write rows to sheet
                    next_row = write_rows_to_sheet(ws, new_rows_data, next_row, source_row)
                    period_added += len(new_rows_data)
                    total_added_all += len(new_rows_data)
                    files_with_data.add(normalized_excel_file)  # Mark this file as having data
                    # Save after each period's data is written
                    print(f"DEBUG: Sauvegarde du fichier cible après période {period}: {normalized_target}")
                    wb.save(normalized_target)
            
            if period_added > 0:
                periods_with_data.append(period)
                periods_processed.append(f"Period {period}: {period_added} records")
        
        # Rename files that had data transferred (only rename once per file)
        for excel_file in files_with_data:
            # Normalize file path
            normalized_excel_file = os.path.normpath(excel_file)
            
            # Double-check file hasn't been renamed already
            if os.path.exists(normalized_excel_file) and not os.path.basename(normalized_excel_file).startswith("(TRANSFER COMPLETED)-"):
                try:
                    file_dir = os.path.dirname(normalized_excel_file)
                    file_name = os.path.basename(normalized_excel_file)
                    new_name = f"(TRANSFER COMPLETED)-{file_name}"
                    new_path = os.path.normpath(os.path.join(file_dir, new_name))
                    print(f"DEBUG: Tentative de renommage: {normalized_excel_file} -> {new_path}")
                    
                    # If a file with the new name already exists, add a number suffix
                    counter = 1
                    original_new_path = new_path
                    while os.path.exists(new_path):
                        name_without_ext, ext = os.path.splitext(original_new_path)
                        new_path = f"{name_without_ext} ({counter}){ext}"
                        counter += 1
                    
                    os.rename(normalized_excel_file, new_path)
                    print(f"DEBUG: Fichier renommé avec succès: {new_path}")
                except Exception as rename_error:
                    error_msg = f"Warning: Could not rename {os.path.basename(normalized_excel_file)}: {str(rename_error)}"
                    print(f"DEBUG: ERREUR lors du renommage - {error_msg}")
                    all_warnings.append(error_msg)
        
        # Build result message
        result_parts = [
            f"Processed all periods (1-26).",
            f"Found data in {len(periods_with_data)} period(s): {', '.join(map(str, periods_with_data))}",
            f"Total records added: {total_added_all}"
        ]
        
        if periods_processed:
            result_parts.append("\nBreakdown by period:")
            result_parts.extend(periods_processed)
        
        if all_warnings:
            result_parts.append(f"\nWarnings:\n" + "\n".join(all_warnings))
        
        if total_added_all == 0:
            return "No records found for any period in the files."
        
        return "\n".join(result_parts)

    except FileNotFoundError as e:
        return f"Error: File not found - {str(e)}"
    except ValueError as e:
        if "Worksheet named 'CNESST' not found" in str(e) or "Worksheet named 'Suivi' not found" in str(e):
            return f"Error: Required worksheet not found. Please ensure files contain 'CNESST' and 'Suivi' sheets."
        return f"Error: {str(e)}"
    except Exception as e:
        return f"Error: {str(e)}"
